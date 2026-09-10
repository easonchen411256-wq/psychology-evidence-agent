"""python-docx infrastructure adapter for manual-review report export."""

from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..literature_review_report import (
    access_label,
    evidence_level_label,
    text,
    truncate,
)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CONTENT_WIDTH_DXA = 9360
DocumentHandle = Any


def set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell: Any, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths_dxa: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)


def set_cell_text(
    cell: Any, value: Any, *, bold: bool = False, color: RGBColor | None = None, size: int = 9
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text(value, ""))
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def configure_document(document: DocumentHandle) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Psychology Evidence Agent | 文献人工审查")
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("仅供人工审查；摘要与初筛不替代全文核对")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED


def add_title(document: DocumentHandle, report: dict[str, Any], counts: dict[str, int]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("文献人工审查报告")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(23)
    run.font.color.rgb = DARK_BLUE
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run(text(report.get("topic"), "心理学证据检索结果"))
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = MUTED
    document.add_heading("研究问题", level=1)
    document.add_paragraph(text(report.get("research_question"), "未在检索报告中记录研究问题。"))
    document.add_heading("本次审查概览", level=1)
    table = document.add_table(rows=3, cols=2)
    set_table_geometry(table, [2700, 6660])
    overview = [
        ("检索候选数", counts["candidate_count"]),
        ("已初筛数", counts["screened_count"]),
        ("优先阅读数", counts["priority_count"]),
        ("本报告收录", counts["selected_count"]),
        ("已发现开放 PDF", counts["open_pdf_count"]),
        ("待人工获取全文", counts["manual_queue_count"]),
    ]
    for index, (label, value) in enumerate(overview):
        row, column = divmod(index, 2)
        set_cell_shading(table.cell(row, column), LIGHT_BLUE)
        set_cell_text(
            table.cell(row, column), f"{label}：{value}", bold=True, color=DARK_BLUE, size=10
        )
    document.add_paragraph(
        "阅读提示：优先查看“初筛理由”和“人工核对提示”，再结合摘要决定是否获取全文。相邻证据或背景证据不能自动视为直接干预效果。"
    )


def add_review_index(document: DocumentHandle, records: list[dict[str, Any]]) -> None:
    document.add_heading("审查索引", level=1)
    document.add_paragraph("建议先在此页填写初步决定，再阅读后续的逐篇阅读卡。")
    table = document.add_table(rows=1, cols=6)
    widths = [600, 3400, 620, 700, 1450, 2590]
    set_table_geometry(table, widths)
    headers = ["序号", "题名", "年份", "得分", "证据层级", "全文状态 / 人工决定"]
    for index, header in enumerate(headers):
        set_cell_shading(table.cell(0, index), LIGHT_BLUE)
        set_cell_text(table.cell(0, index), header, bold=True, color=DARK_BLUE, size=9)
    for number, record in enumerate(records, start=1):
        cells = table.add_row().cells
        values = [
            number,
            truncate(record.get("title"), 150),
            record.get("year", ""),
            record.get("relevance_score", ""),
            evidence_level_label(record.get("evidence_level")),
            f"{access_label(record.get('access'))}\n决定：□保留 □排除 □待定",
        ]
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, size=8)


def add_detail_card(
    document: DocumentHandle, record: dict[str, Any], number: int, abstract_limit: int
) -> None:
    document.add_page_break()
    document.add_heading(f"{number}. {text(record.get('title'), '题名未报告')}", level=1)
    table = document.add_table(rows=5, cols=2)
    set_table_geometry(table, [1800, 7560])
    metadata = [
        ("年份 / 期刊", f"{text(record.get('year'))} / {text(record.get('venue'))}"),
        (
            "作者",
            "; ".join(record.get("authors", []))
            if isinstance(record.get("authors"), list)
            else text(record.get("authors")),
        ),
        ("DOI", text(record.get("doi"))),
        (
            "初筛判断",
            f"得分：{text(record.get('relevance_score'))}；{evidence_level_label(record.get('evidence_level'))}；主题：{text(record.get('subtopic'))}",
        ),
        ("全文状态", access_label(record.get("access"))),
    ]
    for row, (label, value) in enumerate(metadata):
        set_cell_shading(table.cell(row, 0), LIGHT_BLUE)
        set_cell_text(table.cell(row, 0), label, bold=True, color=DARK_BLUE, size=9)
        set_cell_text(table.cell(row, 1), value, size=9)
    document.add_heading("摘要", level=2)
    document.add_paragraph(truncate(record.get("abstract"), abstract_limit))
    document.add_heading("初筛理由", level=2)
    document.add_paragraph(text(record.get("rationale"), "未提供初筛理由。"))
    document.add_heading("人工核对提示", level=2)
    document.add_paragraph(
        text(record.get("human_review_note"), "请结合全文确认研究设计、指标和结论边界。")
    )
    manual = record.get("manual_retrieval")
    if isinstance(manual, dict):
        document.add_heading("全文获取提示", level=2)
        document.add_paragraph(text(manual.get("reason")))
        pages = manual.get("article_pages")
        if isinstance(pages, list) and pages:
            document.add_paragraph("可人工访问的页面：" + "\n".join(str(page) for page in pages))
    document.add_heading("我的人工审查记录", level=2)
    decision = document.add_table(rows=2, cols=2)
    set_table_geometry(decision, [2200, 7160])
    set_cell_shading(decision.cell(0, 0), LIGHT_GRAY)
    set_cell_shading(decision.cell(1, 0), LIGHT_GRAY)
    set_cell_text(decision.cell(0, 0), "决定", bold=True, color=DARK_BLUE, size=10)
    set_cell_text(decision.cell(0, 1), "□ 保留    □ 排除    □ 待定", size=10)
    set_cell_text(decision.cell(1, 0), "理由 / 下一步", bold=True, color=DARK_BLUE, size=10)
    set_cell_text(decision.cell(1, 1), "", size=10)


class WordReportExporter:
    """Render prepared review data using python-docx.

    The ``Any`` document handle is intentionally confined to this adapter because
    python-docx does not provide complete type information for its object model.
    """

    def export(
        self,
        report: dict[str, Any],
        records: list[dict[str, Any]],
        counts: dict[str, int],
        output_path: Path | BinaryIO,
        *,
        abstract_limit: int,
    ) -> None:
        document: DocumentHandle = Document()
        configure_document(document)
        add_title(document, report, counts)
        add_review_index(document, records)
        for number, record in enumerate(records, start=1):
            add_detail_card(document, record, number, abstract_limit)
        document.core_properties.title = "文献人工审查报告"
        document.core_properties.subject = "Psychology Evidence Agent literature review"
        document.core_properties.comments = "Generated from local search and screening records."
        if isinstance(output_path, Path):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            document.save(str(output_path))
        else:
            document.save(output_path)
